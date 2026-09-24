from database import get_connection
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ============================================================
# CONEXIÓN (centralizada en database.get_connection)
# ============================================================

OUTPUT_FILE = "FracttalIntegration.xlsx"

TABLES = [
    "machinery",
    "machine_meters",
    "horometer_updates",
]


# ============================================================
# OBTENER DATOS
# ============================================================

def get_table_data(cursor, table_name):

    cursor.execute(
        f"SELECT * FROM [{table_name}]"
    )

    columns = [
        column[0]
        for column in cursor.description
    ]

    rows = cursor.fetchall()

    return columns, rows


# ============================================================
# FORMATO GENERAL
# ============================================================

def format_sheet(ws):

    # --------------------------------------------------------
    # Encabezados
    # --------------------------------------------------------

    for cell in ws[1]:

        cell.font = Font(
            bold=True
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # --------------------------------------------------------
    # Congelar encabezado
    # --------------------------------------------------------

    ws.freeze_panes = "A2"

    # --------------------------------------------------------
    # Filtro
    # --------------------------------------------------------

    if ws.max_row >= 1 and ws.max_column >= 1:

        ws.auto_filter.ref = ws.dimensions

    # --------------------------------------------------------
    # Ajustar ancho de columnas
    # --------------------------------------------------------

    for column_cells in ws.columns:

        max_length = 0

        for cell in column_cells:

            value = (
                ""
                if cell.value is None
                else str(cell.value)
            )

            max_length = max(
                max_length,
                len(value)
            )

        width = min(
            max(max_length + 2, 10),
            45
        )

        column_letter = get_column_letter(
            column_cells[0].column
        )

        ws.column_dimensions[
            column_letter
        ].width = width

    # --------------------------------------------------------
    # Formato de fechas
    # --------------------------------------------------------

    for row in ws.iter_rows():

        for cell in row:

            if cell.value is not None:

                if hasattr(
                    cell.value,
                    "year"
                ):

                    cell.number_format = (
                        "dd-mm-yyyy hh:mm:ss"
                    )


# ============================================================
# CREAR TABLA DE EXCEL
# ============================================================

def create_excel_table(ws, table_name):

    if ws.max_row < 2:
        return

    ref = ws.dimensions

    table = Table(
        displayName=table_name,
        ref=ref
    )

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )

    table.tableStyleInfo = style

    ws.add_table(table)


# ============================================================
# CREAR HOJA DE TABLA
# ============================================================

def create_table_sheet(
    workbook,
    cursor,
    table_name
):

    print()
    print(
        f"Exportando tabla: {table_name}..."
    )

    columns, rows = get_table_data(
        cursor,
        table_name
    )

    # ========================================================
    # CREAR HOJA
    # ========================================================

    ws = workbook.create_sheet(
        title=table_name[:31]
    )

    # --------------------------------------------------------
    # ENCABEZADOS
    # --------------------------------------------------------

    ws.append(columns)

    # --------------------------------------------------------
    # REGISTROS
    # --------------------------------------------------------

    for row in rows:

        ws.append(
            list(row)
        )

    # --------------------------------------------------------
    # FORMATO
    # --------------------------------------------------------

    format_sheet(ws)

    # --------------------------------------------------------
    # TABLA EXCEL
    # --------------------------------------------------------

    create_excel_table(
        ws,
        f"tbl_{table_name}"
    )

    print(
        f"[OK] {table_name}: "
        f"{len(rows)} registros | "
        f"{len(columns)} columnas"
    )

    return len(rows)


# ============================================================
# CLASIFICACIÓN DE FLOTA
# ============================================================

def create_classification_sheet(
    workbook,
    cursor
):

    print()
    print("Creando hoja CLASIFICACIÓN...")

    # ========================================================
    # OBTENER CLASIFICACIÓN DIRECTAMENTE DESDE SQL
    # ========================================================

    cursor.execute("""
        SELECT
            ISNULL(asset_type, 'NO CLASIFICADO') AS asset_type,
            COUNT(*) AS cantidad
        FROM machinery
        GROUP BY asset_type
        ORDER BY cantidad DESC, asset_type
    """)

    rows = cursor.fetchall()

    # ========================================================
    # CREAR HOJA
    # ========================================================

    ws = workbook.create_sheet(
        title="CLASIFICACIÓN"
    )

    headers = [
        "asset_type",
        "cantidad",
    ]

    ws.append(headers)

    # ========================================================
    # REGISTROS
    # ========================================================

    total = 0

    for asset_type, count in rows:

        ws.append([
            asset_type,
            count,
        ])

        total += count

    # ========================================================
    # TOTAL
    # ========================================================

    ws.append([
        "TOTAL",
        total,
    ])

    # ========================================================
    # FORMATO
    # ========================================================

    format_sheet(ws)

    # --------------------------------------------------------
    # Destacar total
    # --------------------------------------------------------

    total_row = ws.max_row

    for cell in ws[total_row]:

        cell.font = Font(
            bold=True
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )

    # ========================================================
    # TABLA EXCEL
    # ========================================================

    if ws.max_row > 2:

        table = Table(
            displayName="tbl_clasificacion",
            ref=f"A1:B{ws.max_row - 1}"
        )

        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        table.tableStyleInfo = style

        ws.add_table(table)

    print(
        f"[OK] CLASIFICACIÓN: "
        f"{len(rows)} tipos | "
        f"{total} equipos"
    )


# ============================================================
# RESUMEN
# ============================================================

def create_summary_sheet(
    workbook,
    cursor
):

    ws = workbook.create_sheet(
        title="RESUMEN",
        index=0
    )

    # ========================================================
    # TÍTULO
    # ========================================================

    ws["A1"] = "FRACTTAL INTEGRATION"

    ws["A1"].font = Font(
        bold=True,
        size=18
    )

    ws.merge_cells(
        "A1:B1"
    )

    # ========================================================
    # ENCABEZADOS
    # ========================================================

    ws["A3"] = "INDICADOR"
    ws["B3"] = "VALOR"

    for cell in ws[3]:

        cell.font = Font(
            bold=True
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # ========================================================
    # CANTIDAD DE MÁQUINAS
    # ========================================================

    cursor.execute("""
        SELECT COUNT(*)
        FROM machinery
    """)

    machinery_count = cursor.fetchone()[0]

    # ========================================================
    # CANTIDAD DE HORÓMETROS
    # ========================================================

    cursor.execute("""
        SELECT COUNT(*)
        FROM machine_meters
    """)

    meter_count = cursor.fetchone()[0]

    # ========================================================
    # CANTIDAD DE ACTUALIZACIONES
    # ========================================================

    cursor.execute("""
        SELECT COUNT(*)
        FROM horometer_updates
    """)

    updates_count = cursor.fetchone()[0]

    # ========================================================
    # ESTADOS
    # ========================================================

    cursor.execute("""
        SELECT
            status,
            COUNT(*) AS cantidad
        FROM horometer_updates
        GROUP BY status
    """)

    status_rows = cursor.fetchall()

    status_counts = {}

    for status, count in status_rows:

        status_counts[
            str(status)
        ] = count

    update_count = status_counts.get(
        "UPDATE",
        0
    )

    skip_count = status_counts.get(
        "SKIP",
        0
    )

    reject_count = status_counts.get(
        "REJECT",
        0
    )

    # ========================================================
    # ÚLTIMA ACTUALIZACIÓN
    # ========================================================

    cursor.execute("""
        SELECT MAX(reading_date)
        FROM horometer_updates
    """)

    last_update = cursor.fetchone()[0]

    # ========================================================
    # CANTIDAD DE EQUIPOS CLASIFICADOS
    # ========================================================

    cursor.execute("""
        SELECT COUNT(*)
        FROM machinery
        WHERE asset_type IS NOT NULL
          AND LTRIM(RTRIM(asset_type)) <> ''
    """)

    classified_count = cursor.fetchone()[0]

    # ========================================================
    # CANTIDAD DE EQUIPOS SIN CLASIFICAR
    # ========================================================

    unclassified_count = (
        machinery_count - classified_count
    )

    # ========================================================
    # ESCRIBIR RESUMEN
    # ========================================================

    summary = [
        ("Máquinas", machinery_count),
        ("Máquinas clasificadas", classified_count),
        ("Máquinas sin clasificar", unclassified_count),
        ("Horómetros", meter_count),
        ("Registros de historial", updates_count),
        ("Actualizaciones", update_count),
        ("Sin cambios", skip_count),
        ("Rechazados", reject_count),
        ("Última lectura", last_update),
    ]

    row_number = 4

    for indicator, value in summary:

        ws.cell(
            row=row_number,
            column=1,
            value=indicator
        )

        ws.cell(
            row=row_number,
            column=2,
            value=value
        )

        row_number += 1

    # ========================================================
    # FORMATO
    # ========================================================

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 25

    ws.freeze_panes = "A4"

    # Última lectura
    if last_update is not None:

        # La última lectura está en B12
        ws["B12"].number_format = (
            "dd-mm-yyyy hh:mm:ss"
        )


# ============================================================
# MAIN
# ============================================================

def main():

    print("=" * 70)
    print("EXPORTACIÓN SQL SERVER → EXCEL")
    print("=" * 70)

    print()
    print("Conectando a SQL Server...")

    connection = get_connection()

    print("[OK] Conexión establecida.")

    cursor = None

    try:

        cursor = connection.cursor()

        workbook = Workbook()

        # ====================================================
        # ELIMINAR HOJA INICIAL
        # ====================================================

        default_sheet = workbook.active

        workbook.remove(
            default_sheet
        )

        # ====================================================
        # EXPORTAR TABLAS
        # ====================================================

        for table_name in TABLES:

            try:

                create_table_sheet(
                    workbook,
                    cursor,
                    table_name
                )

            except Exception as e:

                print(
                    f"[ERROR] No se pudo exportar "
                    f"{table_name}: {e}"
                )

        # ====================================================
        # CREAR CLASIFICACIÓN
        # ====================================================

        create_classification_sheet(
            workbook,
            cursor
        )

        # ====================================================
        # CREAR RESUMEN
        # ====================================================

        print()
        print("Creando hoja RESUMEN...")

        create_summary_sheet(
            workbook,
            cursor
        )

        print(
            "[OK] Hoja RESUMEN creada."
        )

        # ====================================================
        # GUARDAR EXCEL
        # ====================================================

        output_path = (
            Path.cwd() / OUTPUT_FILE
        )

        workbook.save(
            output_path
        )

        print()
        print("=" * 70)
        print("EXPORTACIÓN FINALIZADA")
        print("=" * 70)

        print()
        print(
            "Archivo generado:"
        )

        print(
            output_path
        )

        print()
        print(
            "Abre el archivo con Excel "
            "para revisar la información."
        )

    finally:

        if cursor is not None:
            cursor.close()

        connection.close()


# ============================================================
# EJECUCIÓN
# ============================================================

if __name__ == "__main__":
    main()