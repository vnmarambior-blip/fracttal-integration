import mssql_python
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


CONNECTION_STRING = (
    "Server=localhost;"
    "Database=FracttalIntegration;"
    "Trusted_Connection=yes;"
    "TrustServerCertificate=yes;"
)

OUTPUT_FILE = "FracttalIntegration.xlsx"

TABLES = [
    "machinery",
    "machine_meters",
    "horometer_updates",
]


def get_connection():
    return mssql_python.connect(CONNECTION_STRING)


def get_table_data(cursor, table_name):
    cursor.execute(f"SELECT * FROM [{table_name}]")

    columns = [
        column[0]
        for column in cursor.description
    ]

    rows = cursor.fetchall()

    return columns, rows


def format_sheet(ws):
    # Formato de encabezados
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )
        cell.alignment = Alignment(
            horizontal="center"
        )

    # Congelar encabezados
    ws.freeze_panes = "A2"

    # Activar filtros
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions

    # Ajustar ancho de columnas
    for column_cells in ws.columns:

        max_length = 0

        for cell in column_cells:

            value = (
                ""
                if cell.value is None
                else str(cell.value)
            )

            max_length = max(
                max_length,
                len(value)
            )

        width = min(
            max(max_length + 2, 10),
            45
        )

        column_letter = get_column_letter(
            column_cells[0].column
        )

        ws.column_dimensions[
            column_letter
        ].width = width


def main():

    print("=" * 70)
    print("EXPORTACIÓN SQL SERVER → EXCEL")
    print("=" * 70)

    print()
    print("Conectando a SQL Server...")

    connection = get_connection()

    print("[OK] Conexión establecida.")

    try:

        cursor = connection.cursor()

        workbook = Workbook()

        # Eliminar hoja inicial de Excel
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        for table_name in TABLES:

            print()
            print(
                f"Exportando tabla: {table_name}..."
            )

            try:

                columns, rows = get_table_data(
                    cursor,
                    table_name
                )

                # Crear hoja
                ws = workbook.create_sheet(
                    title=table_name[:31]
                )

                # Encabezados
                ws.append(columns)

                # Registros
                for row in rows:
                    ws.append(list(row))

                # Formato
                format_sheet(ws)

                print(
                    f"[OK] {table_name}: "
                    f"{len(rows)} registros | "
                    f"{len(columns)} columnas"
                )

            except Exception as e:

                print(
                    f"[ERROR] No se pudo exportar "
                    f"{table_name}: {e}"
                )

        # Guardar Excel
        output_path = (
            Path.cwd() / OUTPUT_FILE
        )

        workbook.save(output_path)

        print()
        print("=" * 70)
        print("EXPORTACIÓN FINALIZADA")
        print("=" * 70)

        print()
        print(
            f"Archivo generado:"
        )

        print(output_path)

        print()
        print(
            "Abre el archivo con Excel para "
            "revisar las tablas."
        )

    finally:

        cursor.close()
        connection.close()


if __name__ == "__main__":
    main()